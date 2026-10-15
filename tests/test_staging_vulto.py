"""Staging da Vulto (etapa 2): cópia crua do legado no schema legado_vulto, com arquivo, aba e linha
de origem, sem limpeza, e sem as senhas em texto puro."""

import hashlib
import json
import os
import sqlite3
from datetime import datetime

import pytest
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import select, text

from src import models as m
from src.config.settings import VULTO_DADOS
from src.database.connection import SessionLocal
from src.migracao_vulto import staging
from src.migracao_vulto.__main__ import rodar


def _linhas(db, tabela: str) -> list[dict]:
    return [dict(r) for r in db.execute(text(f'SELECT * FROM legado_vulto."{tabela}" ORDER BY _linha')).mappings()]


@pytest.fixture
def db():
    with SessionLocal() as sessao:
        staging.limpar_schema(sessao)
        yield sessao
        sessao.rollback()


# ---------- Nomes e valores ----------


def test_nome_de_coluna_sem_acento_sem_repetir_e_com_reserva():
    usados: set[str] = set()
    assert staging.nome_coluna("Preço (R$)", "col_a", usados) == "preco_r"
    assert staging.nome_coluna("Preço (R$)", "col_b", usados) == "preco_r_2"
    assert staging.nome_coluna(None, "col_c", usados) == "col_c"
    assert staging.nome_coluna("   ", "col_d", usados) == "col_d"
    assert staging.nome_coluna("2024", "col_e", usados) == "c_2024"
    assert staging.nome_tabela("contagem", "SAVASSI (por modelo)") == "contagem__savassi_por_modelo"


def test_valor_guarda_o_texto_como_veio_e_o_tipo_que_o_excel_deu():
    assert staging.valor_texto("  R$ 149,90 ") == ("  R$ 149,90 ", None)
    assert staging.valor_texto(7897969295881) == ("7897969295881", "int")
    assert staging.valor_texto(129.9) == ("129.9", "float")
    assert staging.valor_texto(datetime(2026, 2, 3)) == ("2026-02-03 00:00:00", "datetime")
    assert staging.valor_texto(True) == ("TRUE", "bool")
    assert staging.valor_texto(None) == (None, None)


# ---------- Origens ----------


def test_sqlite_entra_cru_com_rowid_e_sem_as_senhas(db, tmp_path):
    caminho = tmp_path / "vulto.db"
    origem = sqlite3.connect(caminho)
    origem.executescript(
        """
        CREATE TABLE funcionarios (id INTEGER PRIMARY KEY, codigo TEXT, nome TEXT, senha TEXT);
        INSERT INTO funcionarios VALUES (1, '0101', 'Bia', 'segredo'), (7, 'estagio', 'Estagiário', '');
        CREATE TABLE config (chave TEXT PRIMARY KEY, valor TEXT);
        INSERT INTO config VALUES ('valor_ponto', '0.10'), ('senha_master', 'vulto@2023');
        CREATE TABLE vendas (id INTEGER PRIMARY KEY, total TEXT, qtd INTEGER, fator REAL);
        INSERT INTO vendas VALUES (10, 'R$ 149,90', 2, 1.5), (11, '19081', NULL, NULL);
        """
    )
    origem.commit()
    origem.close()

    cargas = staging.carregar_sqlite(db, str(caminho), ("funcionarios", "config", "vendas"))

    funcionarios = _linhas(db, "db__funcionarios")
    assert "senha" not in funcionarios[0]
    assert [(f["_linha"], f["codigo"], f["_senha_preenchida"]) for f in funcionarios] == [(1, "0101", "1"), (7, "estagio", "0")]
    assert {c["chave"]: c["valor"] for c in _linhas(db, "db__config")} == {"valor_ponto": "0.10", "senha_master": "***"}
    vendas = _linhas(db, "db__vendas")
    assert [(v["_arquivo"], v["_aba"], v["id"], v["total"], v["qtd"], v["fator"]) for v in vendas] == [
        ("vulto.db", "vendas", "10", "R$ 149,90", "2", "1.5"),
        ("vulto.db", "vendas", "11", "19081", None, None),
    ]
    assert [(c.tabela, c.linhas_origem, c.carregadas) for c in cargas] == [
        ("db__funcionarios", 2, 2),
        ("db__config", 2, 2),
        ("db__vendas", 2, 2),
    ]
    with pytest.raises(RuntimeError, match="faltam"):
        staging.carregar_sqlite(db, str(caminho), ("vendas", "lojas"))


def test_planilha_respeita_o_cabecalho_deslocado_e_guarda_tipo_e_formatacao(db, tmp_path):
    livro = Workbook()
    folha = livro.active
    folha.title = "Coleção"
    folha["A1"] = "COLEÇÃO INVERNO 2025 - compras"
    folha.append([])
    folha.append(["cod", "Preço (R$)", "Preço (R$)", None, "data"])  # linha 3: cabeçalho
    folha.append(["vlt cam 012 verde oliva p", "R$ 169,90", 79, "anotação solta", datetime(2026, 2, 3)])
    folha.append([])  # linha vazia no meio
    folha.append(["CAM016GP", 149.9, None, None, "02/03/2026"])
    folha["A6"].font = Font(strike=True)
    folha["B6"].fill = PatternFill(fill_type="solid", fgColor="FFFF0000")
    notas = livro.create_sheet("LEIA ISSO")
    notas.append(["1. Cada coleção tem sua aba."])
    caminho = tmp_path / "catalogo.xlsx"
    livro.save(caminho)

    cargas = staging.carregar_planilha(db, str(caminho), "catalogo", {"Coleção": 3, "LEIA ISSO": None})

    colecao = _linhas(db, "catalogo__colecao")
    assert [r["_linha"] for r in colecao] == [4, 6]
    primeira, segunda = colecao
    assert (primeira["cod"], primeira["preco_r"], primeira["preco_r_2"], primeira["col_d"]) == (
        "vlt cam 012 verde oliva p",
        "R$ 169,90",
        "79",
        "anotação solta",
    )
    assert json.loads(primeira["_tipos"]) == {"preco_r_2": "int", "data": "datetime"}
    assert primeira["data"] == "2026-02-03 00:00:00" and segunda["data"] == "02/03/2026"
    assert (primeira["_tachado"], primeira["_cor_fundo"]) == ("0", None)
    assert (segunda["_tachado"], segunda["_cor_fundo"]) == ("1", "FFFF0000")
    carga = cargas[0]
    assert (carga.linha_cabecalho, carga.titulo, carga.linhas_origem, carga.vazias, carga.carregadas) == (
        3,
        "COLEÇÃO INVERNO 2025 - compras",
        3,
        1,
        2,
    )
    assert [c["original"] for c in carga.colunas] == ["cod", "Preço (R$)", "Preço (R$)", None, "data"]
    assert _linhas(db, "catalogo__leia_isso")[0]["col_a"] == "1. Cada coleção tem sua aba."

    # Aba que ninguém mapeou para a carga: alguém precisa decidir o que ela é
    with pytest.raises(RuntimeError, match="não mapeadas"):
        staging.carregar_planilha(db, str(caminho), "catalogo", {"Coleção": 3})


def test_csv_latin1_com_sep_e_mensagens_que_quebram_linha(db, tmp_path):
    conteudo = (
        "sep=;\r\n"
        "id;canal;contato;mensagem;tags\r\n"
        '40001;insta;@nicole.lima;"Oii, veio com defeito ?\nQuero devolução";"troca;tamanho"\r\n'
        "40002;WhatsApp;11948854328;Meus pontos sumiram;\n"
        ";;;;\r\n"
        "40003;email;ana@x.com;ok;extra;campo a mais\r\n"
    )
    caminho = tmp_path / "atendimentos.csv"
    caminho.write_bytes(conteudo.encode("latin-1"))

    carga = staging.carregar_csv(db, str(caminho), "atendimentos_teste")

    linhas = _linhas(db, "atendimentos_teste")
    assert [(r["_linha"], r["id"]) for r in linhas] == [(3, "40001"), (5, "40002"), (7, "40003")]
    assert linhas[0]["mensagem"] == "Oii, veio com defeito ?\nQuero devolução" and linhas[0]["tags"] == "troca;tamanho"
    assert json.loads(linhas[2]["_colunas_extras"]) == ["campo a mais"]
    assert (carga.titulo, carga.linha_cabecalho, carga.carregadas, carga.vazias) == ("sep=;", 2, 3, 1)


# ---------- Carga completa com os dados reais ----------

_SEM_DADOS = not os.path.isdir(VULTO_DADOS)


def _assinaturas() -> dict[str, str]:
    return {
        n: hashlib.sha256(open(os.path.join(VULTO_DADOS, n), "rb").read()).hexdigest() for n in sorted(os.listdir(VULTO_DADOS))
    }


@pytest.mark.skipif(_SEM_DADOS, reason="Pasta VULTO_DADOS com os dados da Vulto não encontrada.")
def test_staging_real_carrega_tudo_bate_com_a_origem_e_nao_toca_no_legado():
    antes = _assinaturas()
    rodar(["staging"])
    rodar(["staging"])  # de novo: mesmo resultado, sem duplicar
    assert _assinaturas() == antes  # o legado é só leitura

    with SessionLocal() as db:
        cargas = {r["tabela"]: r for r in db.execute(text("SELECT * FROM legado_vulto._cargas")).mappings()}
        esperadas = {
            "db__vendas": 18598,
            "db__itens_venda": 31110,
            "db__compras": 14948,
            "db__pontos_mov": 12127,
            "db__estoque": 7130,
            "db__entradas_estoque": 17331,
            "db__funcionarios": 24,
            "clientes__clientes": 2535,
            "clientes__resgates": 596,
            "catalogo__verao_26": 450,
            "atendimentos": 1664,
        }
        assert {t: cargas[t]["linhas_carregadas"] for t in esperadas} == esperadas
        # Todas as tabelas do vulto.db e todas as abas das planilhas estão lá
        abas = sum(len(abas) for _, abas in staging.PLANILHAS.values())
        assert len(cargas) == len(staging.TABELAS_DB) + abas + 1
        for tabela, carga in cargas.items():
            assert db.scalar(text(f'SELECT count(*) FROM legado_vulto."{tabela}"')) == carga["linhas_carregadas"]

        senha = db.scalar(
            text("SELECT count(*) FROM information_schema.columns WHERE table_schema='legado_vulto' AND column_name='senha'")
        )
        assert senha == 0
        assert db.scalar(text("SELECT valor FROM legado_vulto.db__config WHERE chave = 'senha_master'")) == "***"
        ocorrencias = db.scalars(select(m.MigracaoOcorrencia.codigo_problema).where(m.MigracaoOcorrencia.dominio == "staging"))
        assert sorted(ocorrencias) == ["SENHA_MASTER", "SENHA_TEXTO_PURO"]
        # A sub-tabela INVERNO 26 continua dentro da aba Verão 26, na linha em que está no Excel
        assert db.scalar(text("SELECT codigo FROM legado_vulto.catalogo__verao_26 WHERE _linha = 276")) == "MOL014VP"
