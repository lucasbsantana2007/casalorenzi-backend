"""Staging: cópia crua dos dados da Vulto no schema legado_vulto (primeira camada da migração).

Tudo entra como veio, em colunas TEXT e sem limpeza, com três colunas de rastreio:
- _arquivo: o arquivo de origem;
- _aba: a aba da planilha (no vulto.db, a tabela);
- _linha: a linha no arquivo, como aparece no Excel ou no editor (no vulto.db, o rowid).

Nas planilhas parte da informação está na célula e não no texto, então cada linha guarda também:
- _tipos: o tipo que o Excel deu às células que não são texto (ex.: uma data que o Excel
  interpretou ao contrário vem como datetime, e não como o texto digitado);
- _tachado, _cor_fonte e _cor_fundo: a formatação (peças descontinuadas são tachadas ou vermelhas).

legado_vulto._cargas registra cada tabela: origem, linha do cabeçalho, títulos acima dele, os
nomes originais das colunas e a contagem de linhas na origem e no staging.

Exceção, por segurança: as senhas em texto puro de funcionarios não são copiadas (fica só se havia
senha) e a senha_master do config entra mascarada. As duas viram ocorrência.

A cada carga as tabelas do schema são apagadas e recriadas: rodar de novo dá o mesmo resultado.
O legado é só leitura (SQLite aberto com mode=ro; planilhas e CSV só lidos).
"""

import csv
import io
import json
import os
import re
import sqlite3
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime, time

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from sqlalchemy import Column, Integer, MetaData, Table, Text, text
from sqlalchemy.orm import Session

from src.config.settings import VULTO_DADOS
from src.migracao_vulto.relatorio import Resultado
from src.repositories import migracao_repository

DOMINIO = "staging"
SCHEMA = "legado_vulto"

ARQUIVO_DB = "vulto.db"
TABELAS_DB = ("lojas", "compras", "pontos_mov", "config", "vendas", "itens_venda", "estoque", "entradas_estoque", "funcionarios")

# Planilha → (prefixo das tabelas, {aba: linha do cabeçalho}). A linha conta como no Excel (1 = primeira);
# None = aba de texto, sem cabeçalho (colunas col_a, col_b...). Aba nova numa planilha faz a carga
# parar, para alguém decidir o que ela é antes de seguir.
PLANILHAS = {
    "catalogo_vulto_2026_REVISADO_final.xlsx": (
        "catalogo",
        {
            "VERAO 25": 1,
            "inverno25": 3,
            "Verão 26": 1,
            "Atemporais": 1,
            "PLANILHA ANTIGA NAO USAR": 2,
            "LEIA ISSO - REGRAS": None,
        },
    ),
    "clientes_fidelidade_FINAL_v3 (1).xlsx": (
        "clientes",
        {"Clientes": 1, "clientes antigos (NAO APAGAR)": 1, "Resgates": 1, "LEIA ISSO": None, "Plan1": 1},
    ),
    "contagem_inventario_set26.xlsx": (
        "contagem",
        {"PINHEIROS 20-09": 3, "Eldorado": 1, "SAVASSI (por modelo)": 1, "batel": 1, "CD - posicao WMS 28-09": 1},
    ),
    "vendas_batel_MANUAL_marco.xlsx": ("batel_manual", {"MARÇO": 3, "Plan2": None}),
}

ARQUIVO_CSV = "atendimentos_export_whatsapp_insta.csv"
TABELA_CSV = "atendimentos"
CODIFICACAO_CSV = "latin-1"

# Nunca copiar: (tabela, coluna) → coluna de rastreio que diz só se havia valor
COLUNAS_OMITIDAS = {("funcionarios", "senha"): "_senha_preenchida"}
# Valores mascarados: (tabela, coluna-chave, valor da chave, coluna mascarada)
MASCARADOS = (("config", "chave", "senha_master", "valor"),)
MASCARA = "***"

COR_PADRAO = "FF000000"
SEM_PREENCHIMENTO = "00000000"


@dataclass
class Carga:
    """Uma tabela do staging e de onde ela veio."""

    tabela: str
    arquivo: str
    aba: str | None
    linha_cabecalho: int | None
    titulo: str | None
    colunas: list[dict] = field(default_factory=list)  # [{"coluna", "original", "letra" | "tipo"}]
    linhas_origem: int = 0  # linhas de dados na origem (planilha: inclui as vazias)
    carregadas: int = 0
    vazias: int = 0


# ---------- Nomes e valores ----------


def nome_coluna(original, reserva: str, usados: set[str]) -> str:
    """Nome de coluna SQL a partir do cabeçalho (sem acento, minúsculo, _ no lugar do resto).
    Sem cabeçalho, usa a reserva (ex.: col_k). Repetidos ganham _2, _3..."""
    texto = "" if original is None else str(original)
    base = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode().lower()
    base = re.sub(r"[^a-z0-9]+", "_", base).strip("_")[:50].rstrip("_") or reserva
    if base[0].isdigit():
        base = f"c_{base}"
    nome, n = base, 2
    while nome in usados:
        nome, n = f"{base}_{n}", n + 1
    usados.add(nome)
    return nome


def nome_tabela(prefixo: str, aba: str) -> str:
    return f"{prefixo}__{nome_coluna(aba, 'aba', set())}"


def valor_texto(valor) -> tuple[str | None, str | None]:
    """(texto, tipo) de uma célula. O tipo só vem quando a célula não é texto: o que o Excel guardou
    como número, data ou booleano. Texto fica exatamente como está (com espaços)."""
    if valor is None:
        return None, None
    if isinstance(valor, str):
        return valor, None
    if isinstance(valor, bool):
        return ("TRUE" if valor else "FALSE"), "bool"
    if isinstance(valor, datetime):
        return valor.isoformat(sep=" "), "datetime"
    if isinstance(valor, date):
        return valor.isoformat(), "date"
    if isinstance(valor, time):
        return valor.isoformat(), "time"
    if isinstance(valor, int):
        return str(valor), "int"
    if isinstance(valor, float):
        return repr(valor), "float"
    return str(valor), type(valor).__name__


def _cor(cor) -> str | None:
    rgb = getattr(cor, "rgb", None) if cor is not None and getattr(cor, "type", None) == "rgb" else None
    return rgb if isinstance(rgb, str) else None


def formatacao(celulas) -> tuple[str, str | None, str | None]:
    """(_tachado, _cor_fonte, _cor_fundo) de uma linha: tachado se alguma célula é tachada; as cores
    que não são o padrão (fonte preta, sem preenchimento), separadas por vírgula."""
    tachado, fontes, fundos = False, set(), set()
    for celula in celulas:
        fonte = celula.font
        if fonte is not None and fonte.strike:
            tachado = True
        if fonte is not None and (rgb := _cor(fonte.color)) and rgb != COR_PADRAO:
            fontes.add(rgb)
        preenchimento = celula.fill
        if (
            preenchimento is not None
            and preenchimento.fill_type
            and (rgb := _cor(preenchimento.fgColor))
            and rgb != SEM_PREENCHIMENTO
        ):
            fundos.add(rgb)
    return ("1" if tachado else "0"), (",".join(sorted(fontes)) or None), (",".join(sorted(fundos)) or None)


# ---------- Escrita no schema ----------


def _criar_tabela(db: Session, nome: str, colunas: list[str], extras: list[str]) -> Table:
    tabela = Table(
        nome,
        MetaData(schema=SCHEMA),
        Column("_arquivo", Text, nullable=False),
        Column("_aba", Text),
        Column("_linha", Integer, nullable=False),
        *[Column(c, Text) for c in colunas],
        *[Column(c, Text) for c in extras],
    )
    tabela.create(db.connection())
    return tabela


def _inserir(db: Session, tabela: Table, linhas: list[dict]) -> None:
    if linhas:
        db.execute(tabela.insert(), linhas)


def _contar(db: Session, nome: str) -> int:
    return db.scalar(text(f'SELECT count(*) FROM {SCHEMA}."{nome}"'))


def limpar_schema(db: Session) -> None:
    """Apaga todas as tabelas do staging (o schema fica)."""
    db.execute(text(f"CREATE SCHEMA IF NOT EXISTS {SCHEMA}"))
    nomes = db.scalars(text("SELECT table_name FROM information_schema.tables WHERE table_schema = :s"), {"s": SCHEMA})
    for nome in list(nomes):
        db.execute(text(f'DROP TABLE IF EXISTS {SCHEMA}."{nome}" CASCADE'))


# ---------- Leitura das origens ----------


def carregar_sqlite(db: Session, caminho: str, tabelas: tuple[str, ...] = TABELAS_DB) -> list[Carga]:
    """Cada tabela do vulto.db vira db__<tabela>, com o rowid em _linha."""
    arquivo = os.path.basename(caminho)
    origem = sqlite3.connect(f"file:{caminho}?mode=ro", uri=True)
    cargas = []
    try:
        existentes = {nome for (nome,) in origem.execute("SELECT name FROM sqlite_master WHERE type = 'table'")}
        faltando = set(tabelas) - existentes
        if faltando:
            raise RuntimeError(f"Tabelas que faltam em {arquivo}: {', '.join(sorted(faltando))}")
        for nome_origem in tabelas:
            info = origem.execute(f"PRAGMA table_info({nome_origem})").fetchall()
            copiadas = [(c[1], c[2]) for c in info if (nome_origem, c[1]) not in COLUNAS_OMITIDAS]
            omitidas = {c[1]: COLUNAS_OMITIDAS[(nome_origem, c[1])] for c in info if (nome_origem, c[1]) in COLUNAS_OMITIDAS}
            usados = {"_arquivo", "_aba", "_linha", *omitidas.values()}
            nomes = [nome_coluna(c, c, usados) for c, _ in copiadas]
            carga = Carga(f"db__{nome_origem}", arquivo, nome_origem, None, None)
            carga.colunas = [{"coluna": n, "original": c, "tipo": t} for n, (c, t) in zip(nomes, copiadas, strict=True)]
            carga.colunas += [
                {"coluna": rastreio, "original": f"{c} (não copiada)", "tipo": "omitida"} for c, rastreio in omitidas.items()
            ]
            tabela = _criar_tabela(db, carga.tabela, nomes, list(omitidas.values()))

            selecao = ", ".join(["rowid"] + [f'"{c}"' for c, _ in copiadas] + [f'"{c}"' for c in omitidas])
            linhas = []
            for registro in origem.execute(f"SELECT {selecao} FROM {nome_origem} ORDER BY rowid"):
                rowid, valores, sensiveis = registro[0], registro[1 : 1 + len(copiadas)], registro[1 + len(copiadas) :]
                linha = {"_arquivo": arquivo, "_aba": nome_origem, "_linha": rowid}
                for nome, valor in zip(nomes, valores, strict=True):
                    linha[nome] = valor_texto(valor)[0] if not isinstance(valor, bytes) else valor.hex()
                for rastreio, valor in zip(omitidas.values(), sensiveis, strict=True):
                    linha[rastreio] = "1" if valor not in (None, "") else "0"
                for tabela_m, chave, valor_chave, coluna in MASCARADOS:
                    if tabela_m == nome_origem and linha.get(chave) == valor_chave:
                        linha[coluna] = MASCARA
                linhas.append(linha)
            _inserir(db, tabela, linhas)
            carga.linhas_origem = origem.execute(f"SELECT count(*) FROM {nome_origem}").fetchone()[0]
            carga.carregadas = len(linhas)
            cargas.append(carga)
    finally:
        origem.close()
    return cargas


def carregar_planilha(db: Session, caminho: str, prefixo: str, abas: dict[str, int | None]) -> list[Carga]:
    """Cada aba vira <prefixo>__<aba>. Linhas totalmente vazias não entram (contam em vazias)."""
    arquivo = os.path.basename(caminho)
    livro = load_workbook(caminho)  # sem read_only: a formatação das células só vem assim
    novas = set(livro.sheetnames) - set(abas)
    if novas:
        raise RuntimeError(f"Abas não mapeadas em {arquivo}: {', '.join(sorted(novas))}. Decida o que são em staging.PLANILHAS.")
    faltando = set(abas) - set(livro.sheetnames)
    if faltando:
        raise RuntimeError(f"Abas que faltam em {arquivo}: {', '.join(sorted(faltando))}")

    cargas = []
    for aba, linha_cabecalho in abas.items():
        folha = livro[aba]
        largura = folha.max_column
        usados = {"_arquivo", "_aba", "_linha", "_tipos", "_tachado", "_cor_fonte", "_cor_fundo"}
        cabecalho = [folha.cell(linha_cabecalho, c).value for c in range(1, largura + 1)] if linha_cabecalho else [None] * largura
        letras = [get_column_letter(c) for c in range(1, largura + 1)]
        nomes = [nome_coluna(original, f"col_{letra.lower()}", usados) for original, letra in zip(cabecalho, letras, strict=True)]
        acima = range(1, linha_cabecalho) if linha_cabecalho else range(0)
        titulos = [
            str(folha.cell(r, c).value) for r in acima for c in range(1, largura + 1) if folha.cell(r, c).value not in (None, "")
        ]
        carga = Carga(nome_tabela(prefixo, aba), arquivo, aba, linha_cabecalho, " | ".join(titulos) or None)
        carga.colunas = [
            {"coluna": n, "original": None if o is None else str(o), "letra": letra}
            for n, o, letra in zip(nomes, cabecalho, letras, strict=True)
        ]
        tabela = _criar_tabela(db, carga.tabela, nomes, ["_tipos", "_tachado", "_cor_fonte", "_cor_fundo"])

        primeira = (linha_cabecalho or 0) + 1
        linhas = []
        for numero, celulas in enumerate(folha.iter_rows(min_row=primeira, max_col=largura), start=primeira):
            valores = [valor_texto(c.value) for c in celulas]
            if all(v is None for v, _ in valores):
                carga.vazias += 1
                continue
            tachado, cor_fonte, cor_fundo = formatacao(celulas)
            tipos = {nome: tipo for nome, (_, tipo) in zip(nomes, valores, strict=True) if tipo}
            linha = {
                "_arquivo": arquivo,
                "_aba": aba,
                "_linha": numero,
                "_tachado": tachado,
                "_cor_fonte": cor_fonte,
                "_cor_fundo": cor_fundo,
            }
            linha.update({nome: v for nome, (v, _) in zip(nomes, valores, strict=True)})
            linha["_tipos"] = json.dumps(tipos, ensure_ascii=False) if tipos else None
            linhas.append(linha)
        _inserir(db, tabela, linhas)
        carga.linhas_origem = max(folha.max_row - (linha_cabecalho or 0), 0)
        carga.carregadas = len(linhas)
        cargas.append(carga)
    return cargas


def carregar_csv(db: Session, caminho: str, tabela_nome: str = TABELA_CSV, codificacao: str = CODIFICACAO_CSV) -> Carga:
    """Export da ferramenta de atendimento: primeira linha "sep=;", depois o cabeçalho. Mensagens podem
    quebrar linha dentro das aspas: _linha é a linha física onde o registro começa."""
    arquivo = os.path.basename(caminho)
    with open(caminho, "rb") as f:
        conteudo = f.read().decode(codificacao)
    primeira, _, resto = conteudo.partition("\n")
    deslocamento = 1
    separador = ";"
    if primeira.strip().lower().startswith("sep="):
        separador = primeira.strip()[4:5] or ";"
    else:
        resto, deslocamento = conteudo, 0

    leitor = csv.reader(io.StringIO(resto, newline=""), delimiter=separador)
    cabecalho = next(leitor)
    usados = {"_arquivo", "_aba", "_linha", "_colunas_extras"}
    nomes = [nome_coluna(c, f"col_{i + 1}", usados) for i, c in enumerate(cabecalho)]
    carga = Carga(tabela_nome, arquivo, None, deslocamento + 1, primeira.strip() if deslocamento else None)
    carga.colunas = [{"coluna": n, "original": o} for n, o in zip(nomes, cabecalho, strict=True)]
    tabela = _criar_tabela(db, tabela_nome, nomes, ["_colunas_extras"])

    linhas, fim_anterior = [], leitor.line_num
    for registro in leitor:
        inicio = deslocamento + fim_anterior + 1
        fim_anterior = leitor.line_num
        if not any(campo.strip() for campo in registro):
            carga.vazias += 1
            continue
        linha = {"_arquivo": arquivo, "_aba": None, "_linha": inicio}
        linha.update({nome: (registro[i] if i < len(registro) else None) for i, nome in enumerate(nomes)})
        # Registro com mais campos que o cabeçalho: o excedente não se perde
        linha["_colunas_extras"] = json.dumps(registro[len(nomes) :], ensure_ascii=False) if len(registro) > len(nomes) else None
        linhas.append(linha)
    _inserir(db, tabela, linhas)
    carga.carregadas = len(linhas)
    carga.linhas_origem = len(linhas) + carga.vazias
    return carga


# ---------- Carga completa ----------


def _registrar_cargas(db: Session, cargas: list[Carga]) -> None:
    db.execute(
        text(
            f"""CREATE TABLE {SCHEMA}._cargas (
                tabela TEXT PRIMARY KEY, arquivo TEXT NOT NULL, aba TEXT, linha_cabecalho INTEGER, titulo TEXT,
                colunas TEXT NOT NULL, linhas_origem INTEGER NOT NULL, linhas_carregadas INTEGER NOT NULL,
                linhas_vazias INTEGER NOT NULL, carregado_em TIMESTAMPTZ NOT NULL DEFAULT now())"""
        )
    )
    for c in cargas:
        db.execute(
            text(
                f"""INSERT INTO {SCHEMA}._cargas (tabela, arquivo, aba, linha_cabecalho, titulo, colunas, linhas_origem,
                    linhas_carregadas, linhas_vazias)
                    VALUES (:tabela, :arquivo, :aba, :linha_cabecalho, :titulo, :colunas, :linhas_origem, :carregadas, :vazias)"""
            ),
            {
                "tabela": c.tabela,
                "arquivo": c.arquivo,
                "aba": c.aba,
                "linha_cabecalho": c.linha_cabecalho,
                "titulo": c.titulo,
                "colunas": json.dumps(c.colunas, ensure_ascii=False),
                "linhas_origem": c.linhas_origem,
                "carregadas": c.carregadas,
                "vazias": c.vazias,
            },
        )


def _registrar_ocorrencias(db: Session, cargas: list[Carga]) -> None:
    migracao_repository.limpar_ocorrencias(db, DOMINIO)
    funcionarios = next((c for c in cargas if c.tabela == "db__funcionarios"), None)
    if funcionarios is not None:
        migracao_repository.registrar_ocorrencia(
            db,
            codigo_problema="SENHA_TEXTO_PURO",
            dominio=DOMINIO,
            origem=f"{ARQUIVO_DB}:funcionarios.senha",
            descricao="As senhas dos funcionários estão em texto puro no legado. Não foram copiadas nem para o "
            "staging (fica só _senha_preenchida); os funcionários entram sem senha, pelo convite.",
            acao_tomada="DESCARTADO",
            impacto=funcionarios.carregadas,
            impacto_unidade="REGISTROS",
        )
    if any(c.tabela == "db__config" for c in cargas):
        migracao_repository.registrar_ocorrencia(
            db,
            codigo_problema="SENHA_MASTER",
            dominio=DOMINIO,
            origem=f"{ARQUIVO_DB}:config.senha_master",
            id_origem="senha_master",
            descricao="O legado tem uma senha mestra que entra como qualquer funcionário (falha de segurança). "
            "Entra mascarada no staging e não é usada em lugar nenhum da plataforma.",
            acao_tomada="DESCARTADO",
            valor_original=MASCARA,
        )


def carregar_tudo(db: Session, pasta: str) -> list[Carga]:
    """Apaga o staging e carrega todas as origens de novo (sem confirmar)."""
    if not os.path.isdir(pasta):
        raise RuntimeError(f"Pasta dos dados da Vulto não encontrada: {pasta}. Defina VULTO_DADOS no .env.")
    limpar_schema(db)
    cargas = carregar_sqlite(db, os.path.join(pasta, ARQUIVO_DB))
    for arquivo, (prefixo, abas) in PLANILHAS.items():
        cargas += carregar_planilha(db, os.path.join(pasta, arquivo), prefixo, abas)
    cargas.append(carregar_csv(db, os.path.join(pasta, ARQUIVO_CSV)))

    for carga in cargas:
        no_banco = _contar(db, carga.tabela)
        if no_banco != carga.carregadas or carga.carregadas + carga.vazias != carga.linhas_origem:
            raise RuntimeError(
                f"Contagem não bate em {carga.tabela}: origem {carga.linhas_origem}, vazias {carga.vazias}, "
                f"lidas {carga.carregadas}, no staging {no_banco}"
            )
    _registrar_cargas(db, cargas)
    _registrar_ocorrencias(db, cargas)
    return cargas


def formatar_cargas(cargas: list[Carga]) -> list[str]:
    cabecalho = f"{'tabela do staging':<44} {'origem':<52} {'linhas':>7} {'vazias':>7} {'staging':>8}"
    linhas = [cabecalho, "-" * len(cabecalho)]
    for c in cargas:
        origem = f"{c.arquivo}{f' › {c.aba}' if c.aba else ''}"
        linhas.append(f"{SCHEMA + '.' + c.tabela:<44} {origem[:52]:<52} {c.linhas_origem:>7} {c.vazias:>7} {c.carregadas:>8}")
    return linhas


def migrar(db: Session, pasta: str | None = None) -> Resultado:
    cargas = carregar_tudo(db, pasta or VULTO_DADOS)
    resultado = Resultado(DOMINIO, lidos=sum(c.carregadas for c in cargas), carregados=sum(c.carregadas for c in cargas))
    resultado.notas = formatar_cargas(cargas)
    resultado.ocorrencias = migracao_repository.ocorrencias_por_acao(db, DOMINIO)
    return resultado
